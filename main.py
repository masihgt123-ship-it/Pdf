import os
import re
import shutil
import threading
from io import BytesIO

from kivy.app import App
from kivy.clock import Clock
from kivy.core.image import Image as CoreImage
from kivy.core.window import Window
from kivy.graphics import Color, Line, RoundedRectangle
from kivy.metrics import dp
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.gridlayout import GridLayout
from kivy.uix.image import Image
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.screenmanager import ScreenManager, Screen, NoTransition
from kivy.uix.scrollview import ScrollView
from kivy.uix.textinput import TextInput
from kivy.uix.widget import Widget
from kivy.utils import platform

from pypdf import PdfReader, PdfWriter, Transformation
from pypdf.generic import (
    DictionaryObject,
    NameObject,
    DecodedStreamObject,
)

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


MAX_BYTES = 20 * 1024 * 1024
MAX_PAGES = 300

DOCX = (
    'application/vnd.openxmlformats-officedocument.'
    'wordprocessingml.document'
)


def reader_for(data):
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise ValueError(
            'Encrypted PDFs are not supported. '
            'Open an unencrypted copy.'
        )
    if not 0 < len(reader.pages) <= MAX_PAGES:
        raise ValueError('Supported page count: 1 to 300.')
    return reader


def pdf_bytes(writer):
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def clean_text(text):
    return ''.join(
        char for char in text
        if char in '\n\r\t'
        or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD
        or 0x10000 <= ord(char) <= 0x10FFFF
    )


def word_bytes(texts):
    document = Document()

    for index, text in enumerate(texts):
        if index:
            document.add_page_break()

        cleaned = (
            clean_text(text).strip()
            or '[No text found on this page]'
        )

        for line in cleaned.splitlines():
            paragraph = document.add_paragraph()
            rtl = any(
                '\u0590' <= char <= '\u08ff'
                or '\ufb1d' <= char <= '\ufefc'
                for char in line
            )

            if rtl:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                tag = OxmlElement('w:bidi')
                tag.set(qn('w:val'), '1')
                paragraph._p.get_or_add_pPr().append(tag)

            run = paragraph.add_run(line)
            run.font.name = 'Arial'

            if rtl:
                tag = OxmlElement('w:rtl')
                tag.set(qn('w:val'), '1')
                run._r.get_or_add_rPr().append(tag)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def page_numbers(value, count):
    result = []

    for part in value.replace(' ', '').split(','):
        if not re.fullmatch(r'\d+(-\d+)?', part):
            raise ValueError(
                'Use page numbers like 1-3,5 (English digits).'
            )

        ends = [int(item) for item in part.split('-')]
        first, last = ends[0], ends[-1]

        if not 1 <= first <= last <= count:
            raise ValueError(
                'Page range is outside this document.'
            )

        for number in range(first - 1, last):
            if number not in result:
                result.append(number)

    return result


class SignaturePad(Widget):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.strokes = []
        self.active = None
        self.line = None

    def normal(self, touch):
        return (
            max(
                0,
                min(1, (touch.x - self.x) / max(1, self.width)),
            ),
            max(
                0,
                min(1, (touch.y - self.y) / max(1, self.height)),
            ),
        )

    def on_touch_down(self, touch):
        if (
            not self.collide_point(*touch.pos)
            or self.active is not None
        ):
            return super().on_touch_down(touch)

        touch.grab(self)
        self.active = touch.uid
        self.strokes.append([self.normal(touch)])

        with self.canvas:
            Color(0.2, 0.7, 1, 1)
            self.line = Line(points=list(touch.pos), width=2)

        return True

    def on_touch_move(self, touch):
        if (
            touch.grab_current is self
            and touch.uid == self.active
        ):
            self.strokes[-1].append(self.normal(touch))

            if self.line is None:
                with self.canvas:
                    Color(0.2, 0.7, 1, 1)
                    self.line = Line(
                        points=list(touch.pos),
                        width=2,
                    )
            else:
                self.line.points += list(touch.pos)

            return True

        return super().on_touch_move(touch)

    def on_touch_up(self, touch):
        if touch.grab_current is self:
            touch.ungrab(self)
            self.active = None
            self.line = None
            return True

        return super().on_touch_up(touch)

    def clear(self):
        self.canvas.clear()
        self.strokes = [[]] if self.active is not None else []
        self.line = None


class ToolButton(Button):
    def __init__(
        self,
        tint=(0.106, 0.161, 0.169, 1),
        **kwargs
    ):
        super().__init__(**kwargs)

        self.background_normal = ''
        self.background_down = ''
        self.background_color = (0, 0, 0, 0)
        self.color = (0.93, 0.98, 0.95, 1)
        self.font_size = '14sp'
        self.halign = 'center'
        self.valign = 'middle'
        self.tint = tint

        with self.canvas.before:
            self.fill = Color(*tint)
            self.panel = RoundedRectangle(
                pos=self.pos,
                size=self.size,
                radius=[dp(16)],
            )

        self.bind(
            pos=self.refresh,
            size=self.refresh,
            state=self.refresh,
        )

    def refresh(self, *args):
        self.panel.pos = self.pos
        self.panel.size = self.size
        self.text_size = (
            max(1, self.width - dp(12)),
            self.height,
        )

        if self.state == 'down':
            self.fill.rgba = tuple(
                value * 0.90 for value in self.tint[:3]
            ) + (1,)
        else:
            self.fill.rgba = self.tint


class PDFToolsApp(App):
    def build(self):
        self.title = 'PDF Tools'
        self.busy = False
        self.picker_pending = False
        self.bound = False
        self.data = None
        self.page = 0
        self.result = None
        self.buttons = []
        self.pending_save = None

        Window.clearcolor = (0.063, 0.098, 0.11, 1)

        root = BoxLayout(
            orientation='vertical',
            padding=[dp(16), dp(48), dp(16), dp(64)],
            spacing=dp(10),
        )

        header = BoxLayout(
            size_hint_y=None,
            height=dp(52),
        )
        header.add_widget(Label(
            text='[b]PDF Tools[/b]',
            markup=True,
            color=(0.93, 0.98, 0.95, 1),
            font_size='24sp',
            halign='left',
        ))
        header.add_widget(Label(
            text='OFFLINE',
            color=(0.72, 0.96, 0.85, 1),
            font_size='12sp',
            size_hint_x=0.35,
        ))
        root.add_widget(header)

        self.screens = ScreenManager(
            transition=NoTransition()
        )
        home = Screen(name='home')
        scroll = ScrollView(do_scroll_x=False)

        dashboard = BoxLayout(
            orientation='vertical',
            spacing=dp(16),
            size_hint_y=None,
        )
        dashboard.bind(
            minimum_height=dashboard.setter('height')
        )

        dashboard.add_widget(Label(
            text='[b]Your next\ngreat document.[/b]',
            markup=True,
            color=(0.93, 0.98, 0.95, 1),
            size_hint_y=None,
            height=dp(100),
            font_size='29sp',
        ))
        dashboard.add_widget(Label(
            text='A private space. A simpler workflow.',
            color=(0.55, 0.65, 0.62, 1),
            size_hint_y=None,
            height=dp(30),
            font_size='13sp',
        ))

        opening = ToolButton(
            text='GET STARTED\n\n+  Import PDF',
            tint=(0.72, 0.96, 0.85, 1),
            size_hint_y=None,
            height=dp(140),
        )
        opening.color = (0.08, 0.18, 0.15, 1)
        opening.font_size = '22sp'
        opening.bind(on_release=lambda _: self.pick(101))
        self.buttons.append(opening)
        dashboard.add_widget(opening)

        dashboard.add_widget(Label(
            text='YOUR DOCUMENT TOOLKIT',
            color=(0.55, 0.65, 0.62, 1),
            size_hint_y=None,
            height=dp(30),
            font_size='12sp',
        ))

        grid = GridLayout(
            cols=1,
            spacing=dp(10),
            size_hint_y=None,
        )
        grid.bind(minimum_height=grid.setter('height'))

        tools = [
            (
                'PDF to Word  /  Editable text',
                lambda: self.convert(False),
            ),
            (
                'OCR to Word  /  Scanned pages',
                lambda: self.convert(True),
            ),
            (
                'Merge  /  Combine PDFs',
                lambda: self.pick(102)
                if self.require() else None,
            ),
            ('Extract  /  Select pages', self.extract),
            ('Note  /  Add your words', self.note),
            ('Signature  /  Draw by hand', self.signature),
            ('Export  /  Save your result', self.save),
        ]

        for title, action in tools:
            button = ToolButton(
                text=title,
                size_hint_y=None,
                height=dp(72),
            )
            button.bind(
                on_release=lambda _, fn=action: fn()
            )
            self.buttons.append(button)
            grid.add_widget(button)

        dashboard.add_widget(grid)
        dashboard.add_widget(Label(
            text=(
                'Private by design. No upload required.\n'
                'PDF input: up to 20 MB / 300 pages.'
            ),
            color=(0.55, 0.65, 0.62, 1),
            font_size='12sp',
            size_hint_y=None,
            height=dp(60),
            halign='center',
        ))

        scroll.add_widget(dashboard)
        home.add_widget(scroll)
        self.screens.add_widget(home)

        viewer = Screen(name='viewer')
        view = BoxLayout(
            orientation='vertical',
            spacing=dp(10),
        )
        self.preview = Image()
        view.add_widget(self.preview)

        self.counter = Label(
            text='Open a PDF to preview it',
            color=(0.72, 0.85, 0.79, 1),
            size_hint_y=None,
            height=dp(28),
        )
        view.add_widget(self.counter)

        self.add_row(view, [
            ('Previous', lambda: self.navigate(-1)),
            ('Next', lambda: self.navigate(1)),
        ])
        self.add_row(view, [
            ('Open PDF', lambda: self.pick(101)),
            ('Save result', self.save),
        ])

        viewer.add_widget(view)
        self.screens.add_widget(viewer)
        root.add_widget(self.screens)

        self.status = Label(
            text='Welcome. Open a PDF to get started.',
            color=(0.55, 0.65, 0.62, 1),
            font_size='12sp',
            size_hint_y=None,
            height=dp(52),
            halign='center',
        )
        self.status.bind(
            size=lambda widget, size: setattr(
                widget, 'text_size', (size[0], None)
            )
        )
        root.add_widget(self.status)

        self.add_row(root, [
            ('Home', lambda: self.show_screen('home')),
            (
                'Document',
                lambda: self.show_screen('viewer'),
            ),
        ])

        self.local_pdf = os.path.join(
            self.user_data_dir, 'current.pdf'
        )
        self.ocr_dir = os.path.join(
            self.user_data_dir, 'ocr_legacy_v1'
        )

        return root

    def save(self):
        if self.busy or self.picker_pending:
            return

        if self.result is None:
            self.status.text = (
                'No result yet. Use a tool first.'
            )
            return

        self.pending_save = self.result
        self.pick(103)

    def on_stop(self):
        if self.bound:
            from android import activity
            activity.unbind(
                on_activity_result=self.activity_result
            )
            self.bound = False

    def show_screen(self, name):
        self.screens.current = name

    def add_row(self, root, actions):
        row = BoxLayout(
            spacing=dp(8),
            size_hint_y=None,
            height=dp(48),
        )

        for title, action in actions:
            button = ToolButton(text=title)
            button.bind(
                on_release=lambda _, fn=action: fn()
            )
            self.buttons.append(button)
            row.add_widget(button)

        root.add_widget(row)

    def lock(self):
        for button in self.buttons:
            button.disabled = (
                self.busy or self.picker_pending
            )

    def report(self, text):
        Clock.schedule_once(
            lambda dt: setattr(self.status, 'text', text)
        )

    def require(self):
        if self.data is None:
            self.status.text = 'Open a PDF first.'
            return False
        return True

    def job(self, fn):
        if self.busy:
            return

        self.busy = True
        self.lock()
        self.status.text = 'Working...'

        def done(update, error):
            try:
                if error:
                    self.status.text = 'Error: ' + error
                elif update:
                    update()
            except Exception as exc:
                self.status.text = 'Error: ' + str(exc)
            finally:
                self.busy = False
                self.lock()

        def work():
            try:
                update = fn()
                Clock.schedule_once(
                    lambda dt: done(update, None)
                )
            except Exception as exc:
                Clock.schedule_once(
                    lambda dt, msg=str(exc): done(None, msg)
                )

        threading.Thread(
            target=work, daemon=True
        ).start()

    def on_start(self):
        if platform == 'android':
            from jnius import autoclass
            from android import activity

            self.pdf_bridge = autoclass(
                'org.pdftools.tools.PdfBridge'
            )
            activity.bind(
                on_activity_result=self.activity_result
            )
            self.bound = True
        else:
            self.status.text = (
                'This build requires Android.'
            )
            for button in self.buttons:
                button.disabled = True

    def pick(self, request):
        if self.busy or self.picker_pending:
            return

        from android.runnable import run_on_ui_thread
        from jnius import autoclass

        self.picker_pending = True
        self.lock()

        @run_on_ui_thread
        def launch():
            try:
                Intent = autoclass(
                    'android.content.Intent'
                )
                activity = autoclass(
                    'org.kivy.android.PythonActivity'
                ).mActivity

                action = (
                    Intent.ACTION_CREATE_DOCUMENT
                    if request == 103
                    else Intent.ACTION_OPEN_DOCUMENT
                )
                intent = Intent(action)
                intent.addCategory(
                    Intent.CATEGORY_OPENABLE
                )
                intent.setType(
                    self.pending_save[1]
                    if request == 103
                    else 'application/pdf'
                )

                if request == 103:
                    intent.putExtra(
                        Intent.EXTRA_TITLE,
                        self.pending_save[2],
                    )

                activity.startActivityForResult(
                    intent, request
                )
            except Exception as exc:
                Clock.schedule_once(
                    lambda dt, msg=str(exc):
                    self.picker_failed(msg)
                )

        launch()

    def picker_failed(self, message):
        self.picker_pending = False
        self.lock()
        self.status.text = (
            'File picker error: ' + message
        )

    def activity_result(self, request, code, intent):
        if request not in (101, 102, 103):
            return

        uri = (
            intent.getData()
            if code == -1 and intent is not None
            else None
        )
        Clock.schedule_once(
            lambda dt: self.selected(request, uri)
        )

    def open_uri(self, uri, mode):
        from jnius import autoclass

        resolver = autoclass(
            'org.kivy.android.PythonActivity'
        ).mActivity.getContentResolver()

        descriptor = resolver.openFileDescriptor(
            uri, 'r' if mode == 'rb' else 'wt'
        )
        if descriptor is None:
            raise ValueError(
                'Cannot access the selected file.'
            )

        try:
            fd = descriptor.detachFd()
        finally:
            descriptor.close()

        try:
            return os.fdopen(fd, mode)
        except Exception:
            os.close(fd)
            raise

    def selected(self, request, uri):
        self.picker_pending = False
        self.lock()

        if uri is None:
            self.status.text = 'Selection cancelled.'
            return

        def work():
            if request == 103:
                payload = self.pending_save[0]

                with self.open_uri(uri, 'wb') as stream:
                    stream.write(payload)

                return lambda: setattr(
                    self.status,
                    'text',
                    'File saved successfully.',
                )

            with self.open_uri(uri, 'rb') as stream:
                data = stream.read(MAX_BYTES + 1)

            if len(data) > MAX_BYTES:
                raise ValueError(
                    'Maximum input file size: 20 MB.'
                )

            incoming = reader_for(data)

            if request == 102:
                current = reader_for(self.data)

                if (
                    len(current.pages)
                    + len(incoming.pages)
                    > MAX_PAGES
                ):
                    raise ValueError(
                        'Merged PDF cannot exceed 300 pages.'
                    )

                writer = PdfWriter()
                for reader in (current, incoming):
                    for page in reader.pages:
                        writer.add_page(page)

                data = pdf_bytes(writer)

            message = (
                'PDF opened.'
                if request == 101
                else 'Merged. Save result to export.'
            )

            return self.prepare_document(
                data, message, request == 102
            )

        self.job(work)

    def bridge(self):
        return self.pdf_bridge

    def prepare_document(
        self, data, message, export=True, page=0
    ):
        count = len(reader_for(data).pages)
        page = min(page, count - 1)
        pending = self.local_pdf + '.pending'

        try:
            with open(pending, 'wb') as stream:
                stream.write(data)

            texture = self.render(page, pending)
            os.replace(pending, self.local_pdf)
        finally:
            if os.path.exists(pending):
                os.remove(pending)

        def update():
            self.data = data
            self.page = page
            self.count = count
            self.result = (
                (data, 'application/pdf', 'edited.pdf')
                if export else None
            )

            self.set_preview(texture)
            self.counter.text = (
                'Page %d / %d' % (page + 1, count)
            )
            self.status.text = message

        return update

    def render(self, page, source=None):
        path = os.path.join(
            self.user_data_dir, 'preview.png'
        )
        self.bridge().render(
            source or self.local_pdf, page, path
        )

        with open(path, 'rb') as stream:
            return stream.read()

    def set_preview(self, data):
        self.screens.current = 'viewer'
        self.preview.texture = CoreImage(
            BytesIO(data), ext='png'
        ).texture

    def navigate(self, delta):
        if not self.require():
            return

        page = max(
            0, min(self.count - 1, self.page + delta)
        )

        def work():
            image = self.render(page)

            def update():
                self.page = page
                self.set_preview(image)
                self.counter.text = (
                    'Page %d / %d'
                    % (page + 1, self.count)
                )
                self.status.text = 'Ready.'

            return update

        self.job(work)

    def install_languages(self):
        destination = os.path.join(
            self.ocr_dir, 'tessdata'
        )
        os.makedirs(destination, exist_ok=True)

        assets = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            'tessdata',
        )

        for language in ('fas', 'eng'):
            filename = language + '.traineddata'
            source = os.path.join(assets, filename)
            target = os.path.join(
                destination, filename
            )

            if (
                not os.path.isfile(source)
                or os.path.getsize(source) < 100000
            ):
                raise ValueError(
                    'Missing or incomplete bundled OCR model: '
                    + filename
                    + '. Rebuild with the updated workflow.'
                )

            if (
                os.path.isfile(target)
                and os.path.getsize(target)
                == os.path.getsize(source)
            ):
                continue

            temporary = target + '.tmp'

            try:
                shutil.copyfile(source, temporary)

                if (
                    os.path.getsize(temporary)
                    != os.path.getsize(source)
                ):
                    raise ValueError(
                        'Could not install OCR model: '
                        + filename
                    )

                os.replace(temporary, target)
            finally:
                if os.path.exists(temporary):
                    os.remove(temporary)

    def convert(self, ocr):
        if not self.require():
            return

        def work():
            reader = reader_for(self.data)

            if ocr:
                self.install_languages()

            texts = []

            for index, page in enumerate(reader.pages):
                self.report(
                    ('OCR' if ocr else 'Extracting')
                    + ' page %d / %d'
                    % (index + 1, len(reader.pages))
                )

                if ocr:
                    text = self.bridge().recognize(
                        self.local_pdf,
                        index,
                        self.ocr_dir,
                    )
                else:
                    text = page.extract_text() or ''

                texts.append(text)

            if not any(text.strip() for text in texts):
                raise ValueError(
                    'No text found. Try OCR for scanned PDFs.'
                )

            content = word_bytes(texts)
            empty = sum(
                not text.strip() for text in texts
            )

            def update():
                self.result = (
                    content, DOCX, 'converted.docx'
                )
                self.status.text = (
                    'Word ready. Save result. Empty pages: %d'
                    % empty
                )

            return update

        self.job(work)

    def ask(self, title, hint, callback):
        box = BoxLayout(
            orientation='vertical',
            spacing=10,
            padding=10,
        )
        entry = TextInput(
            hint_text=hint,
            multiline=False,
        )
        box.add_widget(entry)

        row = BoxLayout(
            size_hint_y=None,
            height=50,
        )
        box.add_widget(row)

        popup = Popup(
            title=title,
            content=box,
            size_hint=(0.95, 0.45),
        )

        actions = [
            ('Cancel', popup.dismiss),
            (
                'Apply',
                lambda: (
                    popup.dismiss(),
                    callback(entry.text),
                ),
            ),
        ]

        for text, action in actions:
            button = Button(text=text)
            button.bind(
                on_release=lambda _, fn=action: fn()
            )
            row.add_widget(button)

        popup.open()

    def extract(self):
        if self.require():
            self.ask(
                'Extract pages into a new PDF',
                'Example: 1-3,5',
                self.extract_range,
            )

    def extract_range(self, value):
        def work():
            reader = reader_for(self.data)
            writer = PdfWriter()

            for index in page_numbers(
                value, len(reader.pages)
            ):
                writer.add_page(reader.pages[index])

            return self.prepare_document(
                pdf_bytes(writer),
                'Extracted pages. Save result.',
            )

        self.job(work)

    def note(self):
        if self.require():
            self.ask(
                'Visible note at top (English only)',
                'Short English note, up to 80 characters',
                self.add_note,
            )

    def add_note(self, text):
        if (
            not text.strip()
            or len(text) > 80
            or any(
                not 32 <= ord(char) <= 126
                for char in text
            )
        ):
            self.status.text = (
                'Use 1-80 printable English characters '
                'for PDF notes.'
            )
            return

        self.stamp(note=text)

    def signature(self):
        if not self.require():
            return

        box = BoxLayout(
            orientation='vertical',
            spacing=8,
            padding=10,
        )
        box.add_widget(Label(
            text=(
                'Draw here. Placed at bottom-right '
                'of current page.'
            ),
            size_hint_y=None,
            height=45,
        ))

        pad = SignaturePad()
        box.add_widget(pad)

        row = BoxLayout(
            size_hint_y=None,
            height=50,
        )
        box.add_widget(row)

        popup = Popup(
            title=(
                'Handwritten signature '
                '(not a digital certificate)'
            ),
            content=box,
            size_hint=(0.98, 0.8),
        )

        def apply():
            strokes = [
                stroke[:]
                for stroke in pad.strokes
                if len(stroke) > 1
            ]

            if not strokes:
                return

            popup.dismiss()
            self.stamp(strokes=strokes)

        actions = [
            ('Cancel', popup.dismiss),
            ('Clear', pad.clear),
            ('Apply', apply),
        ]

        for text, action in actions:
            button = Button(text=text)
            button.bind(
                on_release=lambda _, fn=action: fn()
            )
            row.add_widget(button)

        popup.open()

    def stamp(self, note=None, strokes=None):
        def work():
            reader = reader_for(self.data)
            page = reader.pages[self.page]
            page.transfer_rotation_to_content()

            left = float(page.mediabox.left)
            bottom = float(page.mediabox.bottom)
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)

            overlay_writer = PdfWriter()
            overlay = overlay_writer.add_blank_page(
                width=width,
                height=height,
            )

            commands = [
                'q',
                '0 0 0 RG',
                '0 0 0 rg',
                '1.5 w',
            ]

            if note:
                escaped = (
                    note.replace('\\', '\\\\')
                    .replace('(', '\\(')
                    .replace(')', '\\)')
                )

                font = DictionaryObject({
                    NameObject('/Type'):
                        NameObject('/Font'),
                    NameObject('/Subtype'):
                        NameObject('/Type1'),
                    NameObject('/BaseFont'):
                        NameObject('/Helvetica'),
                })

                overlay[NameObject('/Resources')] = (
                    DictionaryObject({
